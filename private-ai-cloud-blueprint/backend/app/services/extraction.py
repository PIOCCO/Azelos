import csv
import io
from pathlib import Path

from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader


def extract_text(filename: str, data: bytes) -> list[dict]:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for i, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            pages.append({"page": i, "section": None, "text": text})
        return pages
    if ext == ".docx":
        doc = DocxDocument(io.BytesIO(data))
        return [{"page": None, "section": "body", "text": "\n".join(p.text for p in doc.paragraphs)}]
    if ext == ".txt" or ext == ".md":
        return [{"page": None, "section": "body", "text": data.decode("utf-8", errors="ignore")}]
    if ext == ".csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8", errors="ignore"))))
        return [{"page": None, "section": "csv", "text": "\n".join([", ".join(r) for r in rows])}]
    if ext == ".xlsx":
        wb = load_workbook(io.BytesIO(data), read_only=True)
        parts = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            lines = []
            for row in ws.iter_rows(values_only=True):
                lines.append(", ".join([str(c) if c is not None else "" for c in row]))
            parts.append(f"Sheet {sheet}:\n" + "\n".join(lines))
        return [{"page": None, "section": "xlsx", "text": "\n\n".join(parts)}]
    raise ValueError(f"Unsupported file type: {ext}")
